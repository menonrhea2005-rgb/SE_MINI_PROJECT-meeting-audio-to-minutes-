
from io import BytesIO

from docx import Document


class DocxExporter:

    @staticmethod
    def export(validated_minutes):
        """Export validated meeting minutes as a Word document."""

        document = Document()

        document.add_heading(validated_minutes.title, level=0)

        document.add_heading("Topics Discussed", level=1)

        if validated_minutes.topics:
            for topic in validated_minutes.topics:
                document.add_paragraph(topic, style="List Bullet")
        else:
            document.add_paragraph("No topics recorded.")

        document.add_heading("Decisions Made", level=1)

        if validated_minutes.decisions:
            for decision in validated_minutes.decisions:
                document.add_paragraph(
                    decision, style="List Bullet"
                )
        else:
            document.add_paragraph("No decisions recorded.")

        document.add_heading("Action Items", level=1)

        if validated_minutes.action_items:
            table = document.add_table(rows=1, cols=3)
            table.style = "Table Grid"

            headers = table.rows[0].cells
            headers[0].text = "Task"
            headers[1].text = "Owner"
            headers[2].text = "Deadline"

            for item in validated_minutes.action_items:
                cells = table.add_row().cells
                cells[0].text = item.task
                cells[1].text = item.owner
                cells[2].text = item.deadline
        else:
            document.add_paragraph("No action items recorded.")

        output = BytesIO()
        document.save(output)

        return output.getvalue()
